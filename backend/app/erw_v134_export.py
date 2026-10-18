"""ERW 1.3.4 — native owner-scoped DOCX and PDF export from saved report objects.

Deterministic report content, source hash verification, embedded data-derived
plots, tables and equation text. No scientific verification or publication.
"""
from __future__ import annotations
from io import BytesIO
from typing import Optional
from fastapi import APIRouter, Header, HTTPException
from fastapi.responses import Response
from docx import Document
from docx.shared import Inches, Pt
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle, Image, KeepTogether
from reportlab.lib import colors
from reportlab.lib.styles import getSampleStyleSheet
from reportlab.lib.utils import ImageReader
from xml.sax.saxutils import escape
from .erw_v121 import owner_from_auth
from .erw_v110 import project_for
from .erw_v120 import REPORT,TABLE,valid_source,SourceRef
from .v1220 import get_calculation
from .v510 import content_hash

router=APIRouter(tags=["erw-1.3.4-native-export"])
NOTICE="DRAFT SCIENTIFIC REPORT — Hypothetical modeled outputs are NOT verified net atmospheric CO2 removal. Sources and citations require independent review."

def get_report(project_id,report_id,owner):
    project_for(project_id,owner)
    saved=get_calculation(report_id,owner)
    if saved.get("projectId")!=project_id:raise HTTPException(404,detail="Report not in selected project")
    obj=saved.get("calculationObject",{})
    if obj.get("schema")!=REPORT:raise HTTPException(422,detail="Not a saved report snapshot")
    h=obj.get("calculationObjectHash")
    if not isinstance(h,str) or h!=content_hash({k:v for k,v in obj.items() if k!="calculationObjectHash"}):
        raise HTTPException(409,detail="Report content hash mismatch")
    if obj.get("verifiedCDR") is not False or obj.get("netDurableAtmosphericRemoval") is not None:
        raise HTTPException(409,detail="Unsupported scientific verification status")
    return obj

def source_parts(project_id,owner,insert):
    ref=insert["source"]
    allowed={"table":{TABLE},
      "figure":{"sc.erw.visual-research-package.v0.6.0","sc.erw.equations-maps-diagrams.v0.7.0"},
      "equation":{"sc.erw.equations-maps-diagrams.v0.7.0"}}
    kind=insert["kind"]
    if kind not in allowed:raise HTTPException(422,detail="Invalid insert kind")
    source=valid_source(project_id,owner,SourceRef(**ref),allowed[kind])
    return source

def render_plot(obj):
    """Produce a real chart from saved ERW 0.6 scenario series, not generated demo values."""
    import matplotlib
    matplotlib.use("Agg")
    import matplotlib.pyplot as plt
    views=obj.get("views") or []
    chart=next((v for v in views if v.get("id")=="scenario-trajectories"),None)
    if not chart: return None
    fig,ax=plt.subplots(figsize=(6.7,3.7))
    try:
        for s in chart.get("series",[]):
            pts=s.get("points",[])
            years=[int(p["year"]) for p in pts]
            values=[float(p["incrementalPotentialProxyTonnesCO2"]) for p in pts]
            if len(years)!=len(values) or not years or any(not 0<=v<1e20 for v in values):
                raise HTTPException(422,detail="Invalid scenario visualization values")
            ax.plot(years,values,marker="o",label=str(s["rateTonnesPerHa"])+" t/ha")
        ax.set_xlabel("Study year")
        ax.set_ylabel("Modeled CO2 potential proxy (tonnes)")
        ax.set_title("Illustrative scenario trajectories — not verified CDR")
        ax.legend(loc="best")
        fig.tight_layout()
        b=BytesIO();fig.savefig(b,format="png",dpi=160);b.seek(0);return b
    finally:plt.close(fig)

def elements(report,project_id,owner):
    inserts=report.get("inserts",[])
    for sec in report.get("sections",[]):
        yield ("heading",sec.get("heading","Section"))
        yield ("body",sec.get("text",""))
        for ins in [x for x in inserts if x["sectionId"]==sec["id"]]:
            obj=source_parts(project_id,owner,ins)
            yield ("caption",str(ins.get("caption") or ins["kind"])+
                   " | source hash: "+str(obj["calculationObjectHash"])+
                   " | illustrative/unverified")
            if ins["kind"]=="table":
                yield ("table",[obj["columns"]]+obj["rows"])
            elif ins["kind"]=="equation":
                for equation in obj.get("equations",[]):
                    yield ("equation",str(equation.get("label",""))+": "+str(equation.get("latex","")))
            else:
                plot=render_plot(obj)
                if plot is not None:yield ("image",plot)
                elif obj.get("equations"):
                    for eq in obj["equations"]:
                        yield ("equation",str(eq.get("label",""))+": "+str(eq.get("latex","")))
                else:yield ("body","Visualization object has no native renderer in this release; refer to saved source hash.")
    yield ("heading","Bibliography")
    yield ("body",report.get("bibliography",""))

def make_docx(report,project_id,owner):
    doc=Document()
    sec=doc.sections[0]
    sec.header.paragraphs[0].text="SUSTAINABLE CATALYST  /  RESEARCH WORKSPACE"
    doc.add_heading(report["title"],0)
    doc.add_paragraph(", ".join(report.get("authors",[])))
    doc.add_paragraph(NOTICE)
    for kind,value in elements(report,project_id,owner):
        if kind=="heading":doc.add_heading(value,level=1)
        elif kind=="body":
            for paragraph in value.split("\n\n"):
                if paragraph.strip():doc.add_paragraph(paragraph)
        elif kind=="caption":doc.add_paragraph(value,style="Caption")
        elif kind=="equation":doc.add_paragraph(value)
        elif kind=="table":
            t=doc.add_table(rows=1,cols=len(value[0]));t.style="Table Grid"
            for idx,cell in enumerate(value[0]):t.rows[0].cells[idx].text=str(cell)
            for row in value[1:]:
                cells=t.add_row().cells
                for i,c in enumerate(row):cells[i].text=str(c)
        elif kind=="image":doc.add_picture(value,width=Inches(6.1))
    b=BytesIO();doc.save(b);return b.getvalue()

def make_pdf(report,project_id,owner):
    styles=getSampleStyleSheet()
    story=[Paragraph(escape(report["title"]),styles["Title"]),
           Paragraph(escape(", ".join(report.get("authors",[]))),styles["Normal"]),
           Spacer(1,12),Paragraph(escape(NOTICE),styles["Italic"]),Spacer(1,16)]
    for kind,value in elements(report,project_id,owner):
        if kind=="heading":story.append(Paragraph(escape(value),styles["Heading2"]))
        elif kind in ("body","equation","caption"):
            style=styles["Normal"] if kind=="body" else styles["Italic"]
            for para in value.split("\n\n"):
                if para.strip():
                    story.append(Paragraph(escape(para).replace("\n","<br/>"),style))
                    story.append(Spacer(1,9))
        elif kind=="table":
            data=[[str(x)[:400] for x in row] for row in value]
            table=Table(data,repeatRows=1,hAlign="LEFT")
            table.setStyle(TableStyle([("BACKGROUND",(0,0),(-1,0),colors.HexColor("#EEEEEE")),
                ("GRID",(0,0),(-1,-1),0.4,colors.grey),
                ("VALIGN",(0,0),(-1,-1),"TOP")]))
            story.append(table);story.append(Spacer(1,12))
        elif kind=="image":
            story.append(Image(value,width=460,height=254));story.append(Spacer(1,10))
    buf=BytesIO()
    SimpleDocTemplate(buf,pagesize=(595.3,841.9),rightMargin=50,leftMargin=50,topMargin=55,bottomMargin=55).build(story)
    return buf.getvalue()

@router.get("/erw/v1.3.4/status")
def status():
    return {"ok":True,"extensionVersion":"1.3.4",
            "nativeDocxExport":True,"nativePdfExport":True,
            "savedSnapshotRequired":True,"ownerScoped":True,
            "hashValidatedSources":True,"figuresFromResearchObjects":True,
            "texEquationTypesetting":False,"completeWordImportRoundtrip":False,
            "publicationEnabled":False,"verifiedCDR":False}

@router.get("/erw/v1.3.4/projects/{project_id}/reports/{report_id}/export/{format}")
def download(project_id:str,report_id:str,format:str,
             authorization:Optional[str]=Header(default=None)):
    owner=owner_from_auth(authorization)
    report=get_report(project_id,report_id,owner)
    if format=="docx":
        blob=make_docx(report,project_id,owner)
        media="application/vnd.openxmlformats-officedocument.wordprocessingml.document"
    elif format=="pdf":
        blob=make_pdf(report,project_id,owner)
        media="application/pdf"
    else:raise HTTPException(404,detail="Unsupported export type")
    return Response(content=blob,media_type=media,headers={
        "Content-Disposition":"attachment; filename=sustainable-catalyst-erw-report."+format,
        "Cache-Control":"private, no-store","X-Content-Type-Options":"nosniff"})
