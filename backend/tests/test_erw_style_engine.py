from io import BytesIO
from zipfile import ZipFile
from docx import Document
from app import erw_style_engine as e
from app import erw_style_ui as ui

REPORT={"title":"Wexford assessment","authors":["Researcher"],
 "sections":[{"id":"abstract","heading":"Abstract","text":"Modelled values are hypothetical."}],
 "inserts":[],"bibliography":"Bibliographic source."}

def test_modern_docx_layout():
 s=e.StyleOptions()
 blob=e.docx_export(REPORT,"project","owner",s)
 d=Document(BytesIO(blob))
 assert abs(d.sections[0].page_height.mm - 297) < 0.1
 assert round(d.sections[0].left_margin.mm)==25
 assert d.styles["Normal"].font.name=="Cambria"
 assert d.styles["Normal"].font.size.pt==11
 assert d.styles["Normal"].paragraph_format.line_spacing==1.5

def test_institutional_word_and_pdf():
 s=e.StyleOptions(preset="institutional")
 word=e.docx_export(REPORT,"project","owner",s)
 d=Document(BytesIO(word))
 assert "SUSTAINABLE CATALYST" in d.sections[0].header.paragraphs[0].text
 pdf=e.pdf_export(REPORT,"project","owner",s)
 assert pdf.startswith(b"%PDF-")

def test_configurable_academic_pdf():
 s=e.StyleOptions(preset="classic-academic",fontFamily="Times New Roman",
                  fontSizePt=12,lineSpacing=2,marginMm=30)
 assert e.pdf_export(REPORT,"project","owner",s).startswith(b"%PDF-")

def test_interface_and_scientific_limitations():
 html=ui.build_page()
 assert 'id="style-preset"' in html
 assert "institutional" in html
 assert "style-docx" in html and "style-pdf" in html
 assert e.status()["verifiedCDR"] is False
 assert e.status()["publicationCertified"] is False
